from pathlib import Path
from shutil import copyfile
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC = Path(r"C:\Users\Administrator\Downloads\ĐỀ CƯƠNG ÔN TẬP GIỮA HK1 NH 2026-2027.docx")
OUT = Path(r"D:\!Download\banh-trang-nuong-indie-upgrade\2026-10-27\l-m-b\outputs\Bo_on_luyen_trac_nghiem_dung_sai_Vat_li_12.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)
copyfile(SRC, OUT)
doc = Document(OUT)

for _style in ('Title', 'Heading 1', 'Heading 2'):
    if _style not in [s.name for s in doc.styles]:
        doc.styles.add_style(_style, WD_STYLE_TYPE.PARAGRAPH)

def set_font(run, size=12, bold=False, color=None):
    run.font.name = 'Arial'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:cs'), 'Arial')
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def border_table(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    tblPr.append(borders)

def para(text='', style=None, size=12, bold=False, before=0, after=5, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align is not None: p.alignment = align
    r = p.add_run(text); set_font(r, size, bold)
    return p

def question(n, prompt, options):
    p = para(f'Câu {n}. {prompt}', size=12, bold=True, before=6, after=2)
    for letter, text in zip('ABCD', options):
        para(f'{letter}. {text}', size=11, after=1)

doc.add_page_break()
p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run('BỘ ÔN LUYỆN TRẮC NGHIỆM VÀ ĐÚNG SAI VẬT LÍ 12')
set_font(r, 18, True)
para('Chủ đề giữa kì 1 năm học 2026 2027', size=13, bold=True, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
para('Bộ câu hỏi này bám theo các nội dung: cấu trúc chất, chuyển thể, nội năng và nguyên lí I nhiệt động lực học, nhiệt độ, nhiệt dung riêng, nhiệt nóng chảy riêng và nhiệt hóa hơi riêng. Hãy làm phần câu hỏi trước, sau đó đối chiếu đáp án ở cuối tài liệu.', size=11, after=10)

para('PHẦN A  TRẮC NGHIỆM NHIỀU LỰA CHỌN', style='Heading 1', size=14, bold=True, before=8, after=5)
mcq = [
('Trong chất rắn, các phân tử thường', ['chuyển động hỗn loạn khắp bình chứa.', 'dao động quanh vị trí cân bằng xác định.', 'đứng yên hoàn toàn.', 'ở rất xa nhau.']),
('Tính chất nào là đặc trưng của chất khí?', ['Có hình dạng và thể tích riêng.', 'Có thể tích riêng nhưng không có hình dạng riêng.', 'Không có hình dạng và thể tích riêng, dễ nén.', 'Các phân tử chỉ dao động quanh vị trí cân bằng.']),
('Hiện tượng nào là sự bay hơi?', ['Nước đá tan thành nước.', 'Hơi nước đọng thành giọt trên thành cốc lạnh.', 'Quần áo ướt khô dần khi phơi.', 'Nước sôi tạo bọt khí.']),
('Sự thăng hoa là quá trình chuyển thể', ['từ rắn sang khí.', 'từ khí sang rắn.', 'từ lỏng sang khí.', 'từ khí sang lỏng.']),
('Ở áp suất xác định, trong khi một chất rắn kết tinh đang nóng chảy, đại lượng nào không đổi?', ['Nội năng.', 'Nhiệt độ.', 'Thể tích.', 'Khối lượng riêng.']),
('Nhiệt độ 27 °C tương ứng gần đúng với', ['246 K.', '273 K.', '300 K.', '310 K.']),
('Điểm đóng băng và điểm sôi của nước ở áp suất chuẩn theo thang Kelvin lần lượt là', ['0 K và 100 K.', '273 K và 373 K.', '32 K và 212 K.', '273 K và 100 K.']),
('Nhiệt dung riêng của một chất cho biết', ['nhiệt lượng để 1 kg chất nóng chảy.', 'nhiệt lượng để 1 kg chất tăng thêm 1 K.', 'nhiệt lượng để toàn bộ chất sôi.', 'nhiệt độ nóng chảy của chất.']),
('Đơn vị SI của nhiệt dung riêng là', ['J/kg.', 'J/K.', 'J/(kg.K).', 'kg.J/K.']),
('Một vật khối lượng m, nhiệt dung riêng c tăng nhiệt độ từ t1 đến t2. Nhiệt lượng vật thu vào là', ['Q = mc(t2 - t1).', 'Q = mc(t2 + t1).', 'Q = m/c(t2 - t1).', 'Q = mc/(t2 - t1).']),
('Nếu hai vật có cùng khối lượng và cùng tăng nhiệt độ 1 K thì vật có nhiệt dung riêng lớn hơn sẽ', ['thu nhiệt lượng lớn hơn.', 'thu nhiệt lượng nhỏ hơn.', 'thu nhiệt lượng bằng nhau.', 'không thu nhiệt lượng.']),
('Nhiệt lượng cần để làm nóng chảy hoàn toàn khối lượng m của chất ở nhiệt độ nóng chảy là', ['Q = mcΔt.', 'Q = mλ.', 'Q = mL.', 'Q = λ/m.']),
('Đơn vị của nhiệt nóng chảy riêng λ là', ['J.', 'J/kg.', 'J/(kg.K).', 'kg/J.']),
('Khi đúc tượng đồng, đồng được nung chảy rồi để nguội trong khuôn. Hai quá trình chuyển thể là', ['nóng chảy và đông đặc.', 'bay hơi và ngưng tụ.', 'thăng hoa và ngưng kết.', 'nóng chảy và bay hơi.']),
('Nhiệt hóa hơi riêng L của nước có ý nghĩa là', ['nhiệt lượng cần để 1 kg nước tăng 1 K.', 'nhiệt lượng cần để 1 kg nước hóa hơi hoàn toàn ở nhiệt độ sôi.', 'nhiệt lượng tỏa ra khi 1 kg nước tăng nhiệt độ.', 'nhiệt lượng cần để làm tan 1 kg nước đá.']),
('Đơn vị của nhiệt hóa hơi riêng là', ['J/kg.', 'J.', 'J/(kg.K).', 'K/J.']),
('Theo quy ước ΔU = A + Q, khi khí nhận nhiệt và đồng thời nở ra sinh công thì', ['Q > 0, A > 0.', 'Q < 0, A < 0.', 'Q > 0, A < 0.', 'Q < 0, A > 0.']),
('Một khí nhận nhiệt lượng 200 J và nhận công 50 J. Độ biến thiên nội năng của khí là', ['150 J.', '200 J.', '250 J.', '-250 J.']),
('Một khí nhận nhiệt lượng 500 J và sinh công 200 J. Độ biến thiên nội năng của khí là', ['700 J.', '300 J.', '-300 J.', '-700 J.']),
('Một khối đồng 2 kg, c = 380 J/(kg.K), tăng nhiệt độ 10 °C. Nhiệt lượng khối đồng thu vào là', ['760 J.', '3 800 J.', '7 600 J.', '76 000 J.']),
('Cần bao nhiêu nhiệt lượng để hóa hơi hoàn toàn 0,20 kg nước ở 100 °C? Cho L = 2,3 x 10^6 J/kg.', ['46 kJ.', '230 kJ.', '460 kJ.', '2 300 kJ.']),
('Một vật có khối lượng 0,5 kg, c = 400 J/(kg.K), nhận 10 000 J. Nhiệt độ của vật tăng thêm', ['5 °C.', '20 °C.', '50 °C.', '80 °C.']),
('Một lượng nước 1 kg được đun từ 20 °C lên 30 °C. Lấy c = 4200 J/(kg.K). Nhiệt lượng nước thu vào là', ['4 200 J.', '21 000 J.', '42 000 J.', '126 000 J.']),
('Quá trình nào làm nội năng của một vật tăng chủ yếu do thực hiện công?', ['Đun nóng nước bằng bếp.', 'Nén nhanh không khí trong bơm xe.', 'Để nước nóng nguội đi.', 'Nước đá tan trong cốc.']),
('Trong nguyên lí I nhiệt động lực học, Q < 0 cho biết', ['hệ nhận nhiệt.', 'hệ truyền nhiệt ra môi trường.', 'hệ nhận công.', 'hệ sinh công.']),
]
for i, (prompt, opts) in enumerate(mcq, 1): question(i, prompt, opts)

doc.add_page_break()
para('PHẦN B  TRẮC NGHIỆM ĐÚNG SAI', style='Heading 1', size=14, bold=True, before=0, after=5)
para('Với mỗi câu, hãy xác định từng phát biểu a, b, c, d là Đúng hay Sai.', size=11, after=6)

tf_sets = [
('Cấu trúc chất và chuyển thể', [
('Chất lỏng có thể tích riêng nhưng không có hình dạng riêng.', 'Đúng'),
('Chất khí rất khó bị nén vì các phân tử ở sát nhau.', 'Sai'),
('Trong quá trình đông đặc của chất tinh khiết, nhiệt độ không đổi.', 'Đúng'),
('Bay hơi chỉ xảy ra ở nhiệt độ sôi.', 'Sai'),
]),
('Nhiệt độ và nhiệt dung riêng', [
('Nhiệt độ 0 °C tương ứng 273 K.', 'Đúng'),
('Khi đổi độ chênh nhiệt độ từ °C sang K thì giá trị độ chênh không đổi.', 'Đúng'),
('Nhiệt dung riêng phụ thuộc vào khối lượng của vật.', 'Sai'),
('Nếu Q = mcΔt thì khi vật tỏa nhiệt, Q có thể mang giá trị âm theo quy ước dấu.', 'Đúng'),
]),
('Nóng chảy và hóa hơi', [
('Công thức tính nhiệt lượng nóng chảy là Q = mλ.', 'Đúng'),
('Nhiệt nóng chảy riêng có đơn vị J/(kg.K).', 'Sai'),
('Để hóa hơi hoàn toàn 1 kg nước ở nhiệt độ sôi cần nhiệt lượng Q = L.', 'Đúng'),
('Khi nước đang sôi ở áp suất không đổi, cung cấp nhiệt làm nhiệt độ nước tiếp tục tăng nhanh.', 'Sai'),
]),
('Nguyên lí I nhiệt động lực học', [
('Theo ΔU = A + Q, nếu khí nhận công thì A > 0.', 'Đúng'),
('Khí nhận nhiệt thì luôn có nội năng tăng.', 'Sai'),
('Nếu khí nhận 100 J nhiệt lượng và sinh công 40 J thì ΔU = 60 J.', 'Đúng'),
('Nếu khí truyền nhiệt ra ngoài và nhận công, Q < 0 và A > 0.', 'Đúng'),
]),
('Bài tập tổng hợp', [
('0,5 kg nhôm có c = 900 J/(kg.K) tăng 20 °C thì thu 9 000 J nhiệt lượng.', 'Đúng'),
('0,1 kg nước hóa hơi hoàn toàn với L = 2,3 x 10^6 J/kg cần 23 000 J.', 'Sai'),
('1 kg nước tăng 5 °C với c = 4200 J/(kg.K) cần 21 000 J.', 'Đúng'),
('Một khí nhận Q = 120 J và nhận A = 80 J thì nội năng giảm 200 J.', 'Sai'),
]),
]
for n, (title, statements) in enumerate(tf_sets, 1):
    para(f'Câu {n}. {title}', size=12, bold=True, before=7, after=3)
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    table.columns[0].width = Cm(13.5); table.columns[1].width = Cm(2.0); table.columns[2].width = Cm(2.0)
    for cell, text in zip(table.rows[0].cells, ['Phát biểu', 'Đúng', 'Sai']):
        cell.text = text; shade(cell, '1F4E78')
        for run in cell.paragraphs[0].runs: set_font(run, 10, True, (255,255,255))
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for idx, (statement, ans) in enumerate(statements):
        cells = table.add_row().cells
        cells[0].text = chr(97+idx) + '. ' + statement
        cells[1].text = '☐'; cells[2].text = '☐'
        for j, cell in enumerate(cells):
            for run in cell.paragraphs[0].runs: set_font(run, 10)
            if j: cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if idx % 2 == 1:
            for cell in cells: shade(cell, 'F3F7FA')
    border_table(table)

doc.add_page_break()
para('ĐÁP ÁN VÀ GỢI Ý TỰ KIỂM TRA', style='Heading 1', size=14, bold=True, before=0, after=5)
para('Phần A', style='Heading 2', size=12, bold=True, before=5, after=3)
answers = ['B','C','C','A','B','C','B','B','C','A','A','B','B','A','B','A','C','C','B','C','C','C','C','B','B']
for start in range(0, len(answers), 5):
    group = '    '.join(f'{i+1}. {answers[i]}' for i in range(start, min(start+5, len(answers))))
    para(group, size=11, after=3)
para('Phần B', style='Heading 2', size=12, bold=True, before=8, after=3)
for n, (_, statements) in enumerate(tf_sets, 1):
    ans = '; '.join(f'{chr(97+i)}. {answer}' for i, (_, answer) in enumerate(statements))
    para(f'Câu {n}: {ans}.', size=11, after=3)
para('Gợi ý: Với các câu tính toán, hãy viết rõ công thức, đổi đơn vị về SI trước khi thay số và kiểm tra đơn vị của kết quả.', size=11, bold=True, before=10, after=0)

for section in doc.sections:
    section.top_margin = Cm(2.0); section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0); section.right_margin = Cm(2.0)

doc.core_properties.title = 'Bộ ôn luyện trắc nghiệm và đúng sai Vật lí 12'
doc.save(OUT)
print(OUT)
